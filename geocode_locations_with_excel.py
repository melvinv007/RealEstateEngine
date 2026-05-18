"""
geocode_locations.py

Reads an Excel file where each row is a hierarchical location path
(e.g. "Dubai>Nad Al Hamar>Building Name"), geocodes rows missing lat/lng
using Nominatim, writes results back to the same Excel file, and optionally
syncs new entries into coordinates.csv.

Usage:
    python geocode_locations.py <path_to_excel.xlsx>
"""

import csv
import os
import sys
import time

from geopy.geocoders import Nominatim
from geopy.exc import GeocoderTimedOut, GeocoderServiceError
from openpyxl import load_workbook

# ── Configuration ──────────────────────────────────────────────────────────────

# Input Excel structure
HIERARCHY_SEPARATOR = ">"           # Separator in the location column
LOCATION_COLUMN_NAME = None         # Column header; None = auto-detect first column

# Output columns written to Excel (columns are created if missing)
LAT_COLUMN_NAME = "lat"
LNG_COLUMN_NAME = "long"

# Geocoding query strategy
GEOCODE_LEVEL = "full"              # "full"  – all segments: "Dubai, Nad Al Hamar, Building"
                                    # "area"  – second segment only: "Nad Al Hamar"
                                    # "last"  – last segment: "Building Name"
                                    # "last2" – last two: "Nad Al Hamar, Building Name"
# APPEND_SUFFIX = "Dubai, UAE"        # Appended to every query; "" to disable
APPEND_SUFFIX = ""        # Appended to every query; "" to disable
SKIP_SUBLOCATIONS = False           # True = skip rows with 3+ hierarchy levels (buildings/POIs)

# Nominatim settings
NOMINATIM_DELAY_SECONDS = 2.0       # Min delay between requests (Nominatim policy: ≥1s)
NOMINATIM_USER_AGENT = "dubai_realestate_geocoder_v1"

# CSV export — sync newly geocoded entries into coordinates.csv
EXPORT_TO_CSV = True
COORDINATES_CSV_PATH = "coordinates.csv"
CSV_CANONICAL_COLUMN = "canonical_name"
CSV_LAT_COLUMN = "lat"
CSV_LNG_COLUMN = "lng"
CSV_CANONICAL_LEVEL = "area"        # Which segment becomes canonical_name in CSV
                                    # Same options as GEOCODE_LEVEL

# Misc
VERBOSE = True      # Print a status line for every row
DRY_RUN = False     # True = geocode and print but write nothing to disk

# ──────────────────────────────────────────────────────────────────────────────


_geolocator = Nominatim(user_agent=NOMINATIM_USER_AGENT, timeout=10)
_last_call: float = 0.0


# ── Nominatim helpers ──────────────────────────────────────────────────────────

def _rate_limit() -> None:
    global _last_call
    elapsed = time.time() - _last_call
    if elapsed < NOMINATIM_DELAY_SECONDS:
        time.sleep(NOMINATIM_DELAY_SECONDS - elapsed)
    _last_call = time.time()


def _geocode(query: str) -> tuple[float, float] | None:
    _rate_limit()
    try:
        result = _geolocator.geocode(query)
        if result:
            return (result.latitude, result.longitude)
        return None
    except GeocoderTimedOut:
        if VERBOSE:
            print(f"\n  [Geocoder] Timeout for '{query}', retrying...")
        time.sleep(NOMINATIM_DELAY_SECONDS * 2)
        try:
            result = _geolocator.geocode(query)
            return (result.latitude, result.longitude) if result else None
        except Exception:
            return None
    except (GeocoderServiceError, Exception) as e:
        if VERBOSE:
            print(f"\n  [Geocoder] Error for '{query}': {e}")
        return None


# ── Path parsing helpers ───────────────────────────────────────────────────────

def _extract_segment(parts: list[str], level: str) -> str:
    """Extract the relevant segment(s) from a split hierarchy path."""
    if level == "full":
        return ", ".join(parts)
    elif level == "area":
        return parts[1] if len(parts) >= 2 else parts[0]
    elif level == "last":
        return parts[-1]
    elif level == "last2":
        return ", ".join(parts[-2:]) if len(parts) >= 2 else parts[-1]
    return parts[-1]


def _build_query(parts: list[str]) -> str:
    base = _extract_segment(parts, GEOCODE_LEVEL)
    return f"{base}, {APPEND_SUFFIX}" if APPEND_SUFFIX else base


def _normalize_header(v: object) -> str:
    return "" if v is None else str(v).strip().lower()


# ── CSV helpers ────────────────────────────────────────────────────────────────

def _load_csv_canonicals(csv_path: str) -> set[str]:
    """Return set of lower-cased canonical_name values already in coordinates.csv."""
    existing: set[str] = set()
    if not os.path.exists(csv_path):
        return existing
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row.get(CSV_CANONICAL_COLUMN, "").strip().lower()
            if name:
                existing.add(name)
    return existing


def _append_to_csv(csv_path: str, canonical: str, lat: float, lng: float) -> None:
    """Append one entry to coordinates.csv, writing headers if the file is new."""
    file_exists = os.path.exists(csv_path)
    with open(csv_path, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow([CSV_CANONICAL_COLUMN, CSV_LAT_COLUMN, CSV_LNG_COLUMN])
        writer.writerow([canonical, lat, lng])


# ── Main ───────────────────────────────────────────────────────────────────────

def geocode_excel(excel_path: str) -> int:
    if not os.path.exists(excel_path):
        print(f"[Error] File not found: {excel_path}")
        return 1

    wb = load_workbook(excel_path)
    ws = wb.active
    rows = list(ws.iter_rows())

    if not rows:
        print("[Error] Excel sheet is empty.")
        return 1

    # ── Identify columns ───────────────────────────────────────────────────────
    header_row = rows[0]
    header_norm = [_normalize_header(cell.value) for cell in header_row]

    # Location column
    if LOCATION_COLUMN_NAME:
        loc_key = _normalize_header(LOCATION_COLUMN_NAME)
        if loc_key not in header_norm:
            print(
                f"[Error] Column '{LOCATION_COLUMN_NAME}' not found. "
                f"Headers found: {[c.value for c in header_row]}"
            )
            return 1
        loc_col_idx = header_norm.index(loc_key)
    else:
        loc_col_idx = 0

    # Lat column — find or create
    lat_key = _normalize_header(LAT_COLUMN_NAME)
    if lat_key in header_norm:
        lat_col_idx = header_norm.index(lat_key)
    else:
        lat_col_idx = len(header_norm)
        ws.cell(row=1, column=lat_col_idx + 1, value=LAT_COLUMN_NAME)
        header_norm.append(lat_key)
        print(f"[Setup] Added column '{LAT_COLUMN_NAME}' at position {lat_col_idx + 1}")

    # Lng column — find or create
    lng_key = _normalize_header(LNG_COLUMN_NAME)
    if lng_key in header_norm:
        lng_col_idx = header_norm.index(lng_key)
    else:
        lng_col_idx = len(header_norm)
        ws.cell(row=1, column=lng_col_idx + 1, value=LNG_COLUMN_NAME)
        header_norm.append(lng_key)
        print(f"[Setup] Added column '{LNG_COLUMN_NAME}' at position {lng_col_idx + 1}")

    # ── Load existing CSV entries ──────────────────────────────────────────────
    csv_existing = _load_csv_canonicals(COORDINATES_CSV_PATH) if EXPORT_TO_CSV else set()

    # ── Process rows ───────────────────────────────────────────────────────────
    stats = {
        "skipped_has_coords": 0,
        "skipped_sublocation": 0,
        "geocoded": 0,
        "failed": 0,
        "csv_added": 0,
    }

    for row_idx, row in enumerate(rows[1:], start=2):
        # Read location value
        loc_cell = row[loc_col_idx] if loc_col_idx < len(row) else None
        raw_value = str(loc_cell.value).strip() if (loc_cell and loc_cell.value) else ""
        if not raw_value:
            continue

        parts = [p.strip() for p in raw_value.split(HIERARCHY_SEPARATOR) if p.strip()]
        if not parts:
            continue

        # Skip sublocations (3+ levels) if configured
        if SKIP_SUBLOCATIONS and len(parts) >= 3:
            stats["skipped_sublocation"] += 1
            if VERBOSE:
                print(f"  Row {row_idx:>4}: SKIP  (sublocation) — {raw_value}")
            continue

        # Skip if both lat and lng are already populated
        def _cell_val(col_idx: int):
            if col_idx < len(row):
                v = row[col_idx].value
                return v if v not in (None, "") else None
            return None

        existing_lat = _cell_val(lat_col_idx)
        existing_lng = _cell_val(lng_col_idx)

        if existing_lat is not None and existing_lng is not None:
            stats["skipped_has_coords"] += 1
            if VERBOSE:
                print(f"  Row {row_idx:>4}: SKIP  (has coords {existing_lat:.4f}, {existing_lng:.4f}) — {raw_value}")
            continue

        # Geocode
        query = _build_query(parts)
        if VERBOSE:
            print(f"  Row {row_idx:>4}: QUERY '{query}' ...", end=" ", flush=True)

        coords = _geocode(query)

        if coords:
            lat, lng = coords
            if VERBOSE:
                print(f"→ ({lat:.5f}, {lng:.5f})")

            if not DRY_RUN:
                ws.cell(row=row_idx, column=lat_col_idx + 1, value=lat)
                ws.cell(row=row_idx, column=lng_col_idx + 1, value=lng)

            stats["geocoded"] += 1

            # CSV sync
            if EXPORT_TO_CSV and not DRY_RUN:
                canonical = _extract_segment(parts, CSV_CANONICAL_LEVEL)
                canonical_lower = canonical.strip().lower()
                if canonical_lower not in csv_existing:
                    _append_to_csv(COORDINATES_CSV_PATH, canonical, lat, lng)
                    csv_existing.add(canonical_lower)
                    stats["csv_added"] += 1
                    if VERBOSE:
                        print(f"           → CSV: added '{canonical}'")
        else:
            if VERBOSE:
                print("→ FAILED")
            stats["failed"] += 1

    # ── Save Excel ─────────────────────────────────────────────────────────────
    if not DRY_RUN:
        wb.save(excel_path)
        print(f"\n[Done] Saved: {excel_path}")
    else:
        print("\n[Dry Run] No files written.")

    print(
        f"[Stats] Geocoded: {stats['geocoded']} | "
        f"Failed: {stats['failed']} | "
        f"Skipped (has coords): {stats['skipped_has_coords']} | "
        f"Skipped (sublocation): {stats['skipped_sublocation']} | "
        f"CSV entries added: {stats['csv_added']}"
    )
    return 0


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python geocode_locations.py <path_to_locations.xlsx>")
        sys.exit(1)
    sys.exit(geocode_excel(sys.argv[1]))


if __name__ == "__main__":
    main()